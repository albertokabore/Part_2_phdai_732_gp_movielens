"""Render the report with native Word bullets, tables, and embedded charts."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt


def render():
    folder = Path(__file__).resolve().parent
    lines = (folder / 'albert_tuning_and_evaluation.md').read_text(encoding='utf-8').splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(1)
    for name in ('Normal', 'List Bullet'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.0

    def inline(paragraph, text):
        for segment in re.split(r'(\*\*.*?\*\*|`[^`]+`|\*[^*]+\*)', text):
            bold = segment.startswith('**')
            italic = segment.startswith('*') and not bold
            code = segment.startswith('`')
            run = paragraph.add_run(segment.strip('*`') if bold or italic or code else segment)
            run.bold, run.italic = bold, italic
            if code:
                run.font.name = 'Consolas'
                run.font.size = Pt(9)

    i, code, appendix = 0, False, False
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith('```'):
            code = not code
        elif code:
            run = doc.add_paragraph().add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        elif line.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[- :|]+\|$', lines[i + 1].strip()):
            headers = [v.strip() for v in line.strip('|').split('|')]
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Light Shading Accent 1'
            for cell, value in zip(table.rows[0].cells, headers):
                inline(cell.paragraphs[0], value)
            i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                for cell, value in zip(table.add_row().cells, lines[i].strip().strip('|').split('|')):
                    inline(cell.paragraphs[0], value.strip())
                i += 1
            continue
        elif line.startswith('#'):
            match = re.match(r'^(#{1,6}) (.*)', line)
            if match:
                title = match.group(2)
                if title.startswith('Appendix:') and not appendix:
                    doc.add_page_break()
                    appendix = True
                doc.add_heading(title, level=min(len(match.group(1)), 4))
        elif re.match(r'^!\[.*\]\(([^)]+)\)$', line):
            target = re.match(r'^!\[.*\]\(([^)]+)\)$', line).group(1)
            doc.add_picture(str((folder / target).resolve()), width=Inches(6.5))
        elif line.startswith('- '):
            inline(doc.add_paragraph(style='List Bullet'), line[2:])
        elif line and line != '---':
            inline(doc.add_paragraph(), line)
        i += 1
    target = folder / 'Albert_Tuning_and_Evaluation_Bullet_Draft.docx'
    doc.save(target)
    print(f'Saved {target.name}: {len(doc.inline_shapes)} images; '
          f'{sum(p.style.name == "List Bullet" for p in doc.paragraphs)} native bullet paragraphs.')
    return target


if __name__ == '__main__':
    render()
